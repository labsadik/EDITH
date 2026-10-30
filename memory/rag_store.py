"""Local document RAG using LangChain splitting, Ollama embeddings and LanceDB."""
from __future__ import annotations

import hashlib
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import requests

BASE_DIR = Path(__file__).resolve().parent.parent
VECTOR_DB_PATH = Path(os.environ.get("EDITH_VECTOR_DB", str(BASE_DIR / "memory" / "vector_db")))
EMBEDDING_URL = os.environ.get("EDITH_EMBEDDING_URL", "http://127.0.0.1:11434").rstrip("/")
EMBEDDING_MODEL = os.environ.get("EDITH_EMBEDDING_MODEL", "nomic-embed-text")
TABLE_NAME = "knowledge_chunks"
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_DOCUMENT_CHARS = 1_000_000


def _dependencies():
    try:
        import lancedb
        from langchain_core.documents import Document
        from langchain_text_splitters import RecursiveCharacterTextSplitter
    except ImportError as exc:
        raise RuntimeError(
            "Local RAG packages are missing. Install requirements.txt "
            "(langchain-core, langchain-text-splitters and lancedb)."
        ) from exc
    return lancedb, Document, RecursiveCharacterTextSplitter


def _read_document(path: Path) -> str:
    if not path.is_file():
        raise FileNotFoundError(f"Knowledge source not found: {path}")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("Knowledge source exceeds the 20 MB indexing limit")
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".rst", ".py", ".json", ".yaml", ".yml", ".csv", ".html", ".xml"}:
        text = path.read_text(encoding="utf-8", errors="replace")
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Install pypdf to index PDF documents") from exc
        text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    elif suffix == ".docx":
        try:
            from docx import Document as WordDocument
        except ImportError as exc:
            raise RuntimeError("Install python-docx to index Word documents") from exc
        text = "\n".join(paragraph.text for paragraph in WordDocument(str(path)).paragraphs)
    else:
        raise ValueError(f"Unsupported RAG file type: {suffix or '(no extension)'}")
    text = text.strip()
    if not text:
        raise ValueError("No extractable text found in the selected document")
    return text[:MAX_DOCUMENT_CHARS]


def _embed(texts: list[str], timeout: float = 90.0) -> list[list[float]]:
    """Create local vectors with Ollama /api/embed, falling back for older servers."""
    if not texts:
        return []
    response = requests.post(
        f"{EMBEDDING_URL}/api/embed",
        json={"model": EMBEDDING_MODEL, "input": texts},
        timeout=timeout,
    )
    if response.status_code in {400, 404}:
        vectors = []
        for value in texts:
            old = requests.post(
                f"{EMBEDDING_URL}/api/embeddings",
                json={"model": EMBEDDING_MODEL, "prompt": value},
                timeout=timeout,
            )
            old.raise_for_status()
            vectors.append(old.json().get("embedding") or [])
    else:
        response.raise_for_status()
        vectors = response.json().get("embeddings") or []
    if len(vectors) != len(texts) or any(not isinstance(v, list) or not v for v in vectors):
        raise RuntimeError(
            f"Ollama embedding model '{EMBEDDING_MODEL}' returned invalid vectors. "
            f"Run: ollama pull {EMBEDDING_MODEL}"
        )
    return [[float(value) for value in vector] for vector in vectors]


def _split_text(text: str, source: str) -> list[dict[str, Any]]:
    _lancedb, Document, Splitter = _dependencies()
    document = Document(page_content=text, metadata={"source": source})
    splitter = Splitter(chunk_size=900, chunk_overlap=140, add_start_index=True)
    chunks = splitter.split_documents([document])[:1200]
    return [
        {
            "text": chunk.page_content,
            "source": source,
            "chunk_id": hashlib.sha256(
                f"{source}\0{chunk.metadata.get('start_index', index)}\0{chunk.page_content}".encode("utf-8")
            ).hexdigest(),
        }
        for index, chunk in enumerate(chunks)
        if chunk.page_content.strip()
    ]


def _table(seed_rows: list[dict[str, Any]] | None = None):
    """Return (table, created_now); seed rows define the real vector dimension."""
    lancedb, _Document, _Splitter = _dependencies()
    VECTOR_DB_PATH.mkdir(parents=True, exist_ok=True)
    db = lancedb.connect(str(VECTOR_DB_PATH))
    try:
        return db.open_table(TABLE_NAME), False
    except Exception:
        if not seed_rows:
            return None, False
        return db.create_table(TABLE_NAME, data=seed_rows), True


def index_text(text: str, source: str = "pasted text") -> dict[str, Any]:
    """Split and embed text, then append chunks to the local vector index."""
    clean = str(text or "").strip()
    if not clean:
        raise ValueError("No text supplied for indexing")
    if len(clean) > MAX_DOCUMENT_CHARS:
        raise ValueError("Text exceeds the 1 MB indexing limit")
    records = _split_text(clean, str(source or "pasted text")[:500])
    if not records:
        raise ValueError("No indexable text chunks were produced")
    vectors = _embed([item["text"] for item in records])
    now = datetime.now(timezone.utc).isoformat()
    rows = [{**item, "vector": vector, "created_at": now}
            for item, vector in zip(records, vectors)]
    table, created = _table(rows)
    if table is None:
        raise RuntimeError("Local vector table could not be created")
    if not created:
        table.add(rows)
    return {
        "ok": True, "source": str(source), "chunks_indexed": len(rows),
        "embedding_model": EMBEDDING_MODEL, "database": str(VECTOR_DB_PATH),
    }


def index_file(path: str | Path) -> dict[str, Any]:
    target = Path(path).expanduser().resolve()
    return index_text(_read_document(target), str(target))


def search_knowledge(query: str, limit: int = 6) -> list[dict[str, Any]]:
    """Return the nearest local chunks and source metadata."""
    question = str(query or "").strip()
    if not question:
        raise ValueError("Search query is required")
    table, _created = _table()
    if table is None:
        return []
    vector = _embed([question])[0]
    rows = table.search(vector).limit(max(1, min(int(limit), 20))).to_list()
    return [
        {"text": str(row.get("text") or ""), "source": str(row.get("source") or ""),
         "score": row.get("_distance")}
        for row in rows if row.get("text")
    ]


def knowledge_stats() -> dict[str, Any]:
    table, _created = _table()
    count = table.count_rows() if table is not None else 0
    return {"database": str(VECTOR_DB_PATH), "table": TABLE_NAME,
            "chunks": int(count), "embedding_model": EMBEDDING_MODEL}
